# -*- coding: utf-8 -*-
"""
影子 AI 决策引擎（Shadow AI）
职责：
  1) 读取 Anomaly_Events.xlsx + Decision_Strategies.xlsx
  2) 接收机载侧上报的实时传感器 JSON（你串口第四部分的归一化格式）
  3) 通过 TRIGGER_RULES 把传感器读数映射到 anomaly_type
  4) 结合当前模式 (CRUISE/PROTECT/MISSION/RTB) 查策略库，
     输出 动作 + Execution_Target + Reward/Penalty + NetScore
设计原则（对应你定的三条底线）：
  - 致命/紧急类的 Execution_Target=FlightController（飞控硬约束兜底）
  - 其余=Shadow_AI（策略级决策）
  - 引擎只读库、不改库；所有评分来自 Decision_Strategies.xlsx
用法：
  from shadow_ai.decision_engine import DecisionEngine
  eng = DecisionEngine()           # 默认读 analyzer/output/ 下两个 xlsx
  eng.set_mode("MISSION")
  decision = eng.decide(sensor_json)
"""
import json, os, sys
from openpyxl import load_workbook
# 把项目根加入 sys.path，让 `import thresholds` 生效（无论从哪里调用）
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)
import thresholds as T   # 阈值集中配置（换电池/电机/抗风包线只改这一个文件）

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EVENTS_XLSX  = os.path.join(BASE, "analyzer","output","Anomaly_Events.xlsx")
STRAT_XLSX   = os.path.join(BASE, "analyzer","output","Decision_Strategies.xlsx")

MODES = ("CRUISE","PROTECT","MISSION","RTB")

# 传感器 metric -> 触发规则 (anomaly_type, 判定函数, 说明)
# 判定函数接收 (value, unit)，返回 bool；数值阈值全部来自 thresholds.py
TRIGGER_RULES = [
    # ---- 内部致命（飞控兜底，权重最高）----
    ("esc_current", "motor_failure",
        lambda v,u: (isinstance(v,(int,float)) and float(v) <= T.MOTOR_FAIL_CURRENT_A),
        f"某电机通道电流<={T.MOTOR_FAIL_CURRENT_A}A => 电机停转/失效"),
    ("vib_high", "prop_break",
        lambda v,u: bool(v),
        "高频振动标志 => 桨叶断裂/不平衡"),
    ("flight_ctrl_alive", "flight_ctrl",
        lambda v,u: (not bool(v)),
        "飞控心跳丢失 => 飞控死机/复位"),
    ("battery_v", "battery_issue",
        lambda v,u: (isinstance(v,(int,float)) and float(v) < T.BATTERY_CRIT_V),
        f"全包电压<{T.BATTERY_CRIT_V}V => 低电量紧急"),
    ("battery_temp", "battery_issue",
        lambda v,u: (isinstance(v,(int,float)) and float(v) > T.BATTERY_TEMP_CRIT_C),
        f"电池温度>{T.BATTERY_TEMP_CRIT_C}°C => 过温"),
    ("esc_temp", "esc_fault",
        lambda v,u: (isinstance(v,(int,float)) and float(v) > T.ESC_TEMP_CRIT_C),
        f"电调温度>{T.ESC_TEMP_CRIT_C}°C => 过热"),
    ("gimbal_ok", "gimbal_fault",
        lambda v,u: (not bool(v)),
        "云台电机不响应 => 云台故障"),
    # ---- 通信/定位 ----
    ("link_lost", "signal_loss",
        lambda v,u: bool(v),
        "通信链路丢失标志 => Failsafe"),
    ("rssi", "signal_loss",
        lambda v,u: (isinstance(v,(int,float)) and float(v) < T.RSSI_LOST_DBM),
        f"RSSI<{T.RSSI_LOST_DBM}dBm => 已丢链"),
    ("hdop", "gps_drift",
        lambda v,u: (isinstance(v,(int,float)) and float(v) > T.GPS_HDOP_MAX),
        f"HDOP>{T.GPS_HDOP_MAX} => GPS 漂移"),
    ("sat_count", "gps_drift",
        lambda v,u: (isinstance(v,(int,float)) and float(v) < T.GPS_SAT_MIN),
        f"卫星数<{T.GPS_SAT_MIN} => 定位降级"),
    # ---- 自然因素 ----
    ("rain", "rain",
        lambda v,u: bool(v),
        "雨滴传感器触发 => 降雨"),
    ("wind_speed", "wind_gust",
        lambda v,u: (isinstance(v,(int,float)) and float(v) > T.WIND_WARN_MS),
        f"风速>{T.WIND_WARN_MS}m/s => 强风告警"),
    ("wind_critical", "wind_gust",
        lambda v,u: (isinstance(v,(int,float)) and float(v) >= T.WIND_CRIT_MS),
        f"风速>={T.WIND_CRIT_MS}m/s => 超抗风包线,强制降落"),
    ("temp_c", "cold_hot",
        lambda v,u: (isinstance(v,(int,float)) and (float(v)<T.TEMP_LOW_C or float(v)>T.TEMP_HIGH_C)),
        f"气温<{T.TEMP_LOW_C}°C 或>{T.TEMP_HIGH_C}°C => 极端温度"),
    ("visibility", "dust_fog",
        lambda v,u: (isinstance(v,(int,float)) and float(v) < T.VISIBILITY_MIN_M),
        f"视距<{T.VISIBILITY_MIN_M}m => 沙尘/雾霾"),
    ("bird_detect", "bird",
        lambda v,u: bool(v),
        "视觉/雷达检出鸟类 => 鸟撞风险"),
    ("thunder", "lightning",
        lambda v,u: bool(v),
        "大气电势/雷击预警 => 立即接地"),
    ("obstacle_close", "obstacle",
        lambda v,u: (isinstance(v,(int,float)) and float(v) < T.OBSTACLE_CLOSE_M),
        f"障碍距离<{T.OBSTACLE_CLOSE_M}m => 接近预警"),
]

class DecisionEngine:
    def __init__(self, events_path=EVENTS_XLSX, strat_path=STRAT_XLSX):
        self.strat = {}     # (anomaly, mode) -> row dict
        self.sev   = {}     # anomaly -> severity
        self._load(strat_path, events_path)

    def _load(self, strat_path, events_path):
        # 异常事件库：取 severity
        wb = load_workbook(events_path); ws = wb["异常事件库"]
        hdr = [c.value for c in ws[1]]
        i_t = hdr.index("Anomaly_Type"); i_s = hdr.index("Severity")
        for r in ws.iter_rows(min_row=2, values_only=True):
            self.sev[r[i_t]] = int(r[i_s] or 0)
        # 策略库
        wb2 = load_workbook(strat_path); ws2 = wb2["决策策略库"]
        h2 = [c.value for c in ws2[1]]
        idx = {name: h2.index(name) for name in h2}
        for r in ws2.iter_rows(min_row=2, values_only=True):
            key = (r[idx["Anomaly_Type"]], r[idx["Mode"]])
            self.strat[key] = {
                "action": r[idx["Action"]],
                "target": r[idx["Execution_Target"]],
                "reward": r[idx["Reward"]],
                "penalty": r[idx["Penalty"]],
                "net": r[idx["NetScore"]],
            }

    def set_mode(self, mode):
        if mode not in MODES:
            raise ValueError(f"mode 必须是 {MODES}")
        self.mode = mode

    def detect_anomalies(self, sensor):
        """sensor: {metric: {value, unit}} 或 {metric: value}
           返回命中的 [(anomaly_type, reason, rule_desc)]"""
        hits = []
        for metric, atype, fn, desc in TRIGGER_RULES:
            if metric not in sensor:
                continue
            val = sensor[metric]
            unit = val.get("unit","") if isinstance(val,dict) else ""
            try:
                fired = fn(val.get("value") if isinstance(val,dict) else val, unit)
            except Exception:
                fired = False
            if fired:
                hits.append((atype, metric, desc))
        return hits

    def decide(self, sensor):
        """主入口：返回 {anomaly, mode, action, target, reward, penalty, net, reasons}"""
        hits = self.detect_anomalies(sensor)
        if not hits:
            return {"anomaly":"none","mode":self.mode,
                    "action":"维持常规巡航","target":"Shadow_AI",
                    "reward":0,"penalty":0,"net":100,"reasons":[]}
        # 多异常时：按 严重度 降序 + 内部因素优先（severity 已编码优先级）
        hits.sort(key=lambda h: -self.sev.get(h[0],0))
        atype = hits[0][0]
        st = self.strat.get((atype, self.mode),
                 {"action":"保守悬停待命","target":"Shadow_AI","reward":0,"penalty":0,"net":50})
        return {
            "anomaly": atype,
            "mode": self.mode,
            "action": st["action"],
            "target": st["target"],
            "reward": int(st["reward"] or 0),
            "penalty": int(st["penalty"] or 0),
            "net": int(st["net"] or 0),
            "severity": self.sev.get(atype,0),
            "reasons": [f"{m}: {d}" for (_a, m, d) in hits],
        }

    def score_model(self, model_answer, ref_answers, threshold=0.8):
        """（可选）与你最初的'相似度奖励机制'对接：
           比较模型回答 vs 参考回答集合，相似度>threshold 奖励，否则惩罚。
           这里用简单 Jaccard 关键词重叠；可替换为 embedding 余弦。"""
        import re
        def toks(t):
            return set(re.findall(r"[\u4e00-\u9fff]{2}|[A-Za-z]{2,}", (t or "").lower()))
        m = toks(model_answer)
        best = 0.0
        for ref in ref_answers:
            r = toks(ref)
            inter = len(m & r); union = len(m | r) or 1
            best = max(best, inter/union)
        return {"similarity": round(best,3),
                "verdict": "reward" if best>=threshold else "penalty",
                "reward": 10 if best>=threshold else 0,
                "penalty": 0 if best>=threshold else 5}

def demo():
    eng = DecisionEngine()
    # 用 thresholds.py 的真实阈值构造示例传感器数据
    samples = [
        ("电机失效(MISSION)", {"esc_current":T.MOTOR_FAIL_CURRENT_A}, "MISSION"),
        ("低电量(PROTECT)", {"battery_v":T.BATTERY_CRIT_V-0.5}, "PROTECT"),
        ("电池过温(MISSION)", {"battery_temp":T.BATTERY_TEMP_CRIT_C+3}, "MISSION"),
        ("降雨(MISSION)", {"rain":1}, "MISSION"),
        ("降雨(PROTECT)", {"rain":1}, "PROTECT"),
        ("大风告警(MISSION)", {"wind_speed":T.WIND_WARN_MS+1}, "MISSION"),
        ("大风超包线(RTB)", {"wind_critical":T.WIND_CRIT_MS+1}, "RTB"),
        ("GPS漂移(RTB)", {"hdop":T.GPS_HDOP_MAX+1.0}, "RTB"),
        ("RSSI弱(PROTECT)", {"rssi":T.RSSI_LOST_DBM-5}, "PROTECT"),
        ("多异常:电低+雨(MISSION)", {"battery_v":T.BATTERY_CRIT_V-0.3,"rain":1}, "MISSION"),
        ("障碍接近(MISSION)", {"obstacle_close":T.OBSTACLE_CLOSE_M-2}, "MISSION"),
        ("雷击预警(RTB)", {"thunder":1}, "RTB"),
    ]
    print(f"{'场景':22s} | anomaly  | target        | R/P     | net | 原因")
    print("-"*95)
    for label, sensor, mode in samples:
        eng.set_mode(mode); d = eng.decide(sensor)
        print(f"{label:22s} | {d['anomaly']:10s} | {d['target']:13s} | "
              f"{d['reward']:>2}/{d['penalty']:<2} | {d['net']:>3} | {d['reasons'][0] if d['reasons'] else '-'}")

if __name__ == "__main__":
    demo()
